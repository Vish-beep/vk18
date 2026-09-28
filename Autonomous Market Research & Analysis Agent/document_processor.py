import os
import re
from io import BytesIO
from pypdf import PdfReader
from docx import Document
from config import CHUNK_SIZE, CHUNK_OVERLAP


def _read_pdf(file_path):
    """Read a PDF defensively and preserve pages that can still be extracted."""
    try:
        size = os.path.getsize(file_path)
    except OSError as exc:
        raise ValueError(f"Could not access PDF: {exc}")

    if size == 0:
        raise ValueError("The uploaded PDF is empty (0 bytes). Upload the complete PDF again.")

    with open(file_path, "rb") as f:
        raw = f.read()

    if not raw.startswith(b"%PDF"):
        raise ValueError("The uploaded file is not a valid PDF file (missing PDF header).")

    # strict=False lets pypdf recover a number of malformed-but-readable PDFs.
    try:
        reader = PdfReader(BytesIO(raw), strict=False)
    except Exception as exc:
        raise ValueError(
            "Could not parse this PDF. The file appears incomplete or corrupted. "
            "Open it in a PDF reader, save/export it as a new PDF, and upload the new file. "
            f"Parser detail: {exc}"
        ) from exc

    if reader.is_encrypted:
        try:
            if reader.decrypt("") == 0:
                raise ValueError("This PDF is password-protected. Remove the password and upload it again.")
        except Exception as exc:
            raise ValueError(f"Could not decrypt the PDF. Remove its password and upload it again. {exc}") from exc

    pages = []
    page_errors = []
    try:
        page_count = len(reader.pages)
    except Exception as exc:
        raise ValueError(f"Could not read the PDF page table: {exc}") from exc

    for page_no in range(page_count):
        try:
            page = reader.pages[page_no]
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[Page {page_no + 1}]\n{text}")
        except Exception as exc:
            page_errors.append(f"page {page_no + 1}: {exc}")
            continue

    result = "\n\n".join(pages)
    if not result.strip():
        detail = ""
        if page_errors:
            detail = " Parser errors: " + "; ".join(page_errors[:3])
        raise ValueError(
            "No readable text could be extracted from this PDF. It may be scanned/image-only "
            "or incomplete/corrupted." + detail
        )

    return result


def extract_text(file_path):
    extension = os.path.splitext(file_path)[1].lower()
    if extension in {".txt", ".md"}:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    if extension == ".pdf":
        return _read_pdf(file_path)
    if extension == ".docx":
        try:
            document = Document(file_path)
        except Exception as exc:
            raise ValueError(f"Could not read DOCX: {exc}") from exc
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    parts.append(row_text)
        result = "\n".join(parts)
        if not result.strip():
            raise ValueError("No extractable text found in DOCX.")
        return result
    raise ValueError(f"Unsupported extension '{extension}'. Use PDF, DOCX, TXT or MD.")


def clean_text(text):
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    if chunk_size <= overlap:
        raise ValueError("chunk_size must be larger than overlap.")
    words = text.split()
    chunks = []
    step = chunk_size - overlap
    for start in range(0, len(words), step):
        chunk = " ".join(words[start:start + chunk_size]).strip()
        if chunk:
            chunks.append(chunk)
        if start + chunk_size >= len(words):
            break
    return chunks
